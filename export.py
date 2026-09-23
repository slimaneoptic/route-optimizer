"""Excel workbook: a summary sheet plus one printable sheet per van, with Google Maps links."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from solver import Plan, Settings, hhmm, maps_legs, plan_cost

HEAD = PatternFill("solid", fgColor="1B3A57")
WHITE = Font(color="FFFFFF", bold=True)


def _style(ws, widths: dict[int, int]) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(vertical="center")
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"


def workbook(stops: pd.DataFrame, plan: Plan, s: Settings, currency: str) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        summary = pd.DataFrame(
            [
                {
                    "van": f"Van {r.vehicle}",
                    "stops": len(r.nodes),
                    "km": round(r.distance_km, 1),
                    "load": r.load,
                    "leaves depot": hhmm(r.start),
                    "back at depot": hhmm(r.end),
                    f"cost ({currency})": round(r.distance_km * s.cost_per_km + s.cost_per_van, 2),
                }
                for r in plan.routes
            ]
        )
        summary.loc[len(summary)] = ["Total", plan.served, round(plan.total_km, 1), summary["load"].sum(), "", "", round(plan_cost(plan, s), 2)]
        summary.to_excel(xl, sheet_name="Summary", index=False)
        ws = xl.sheets["Summary"]
        _style(ws, {1: 10, 2: 8, 3: 9, 4: 8, 5: 14, 6: 15, 7: 13})
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)

        for r in plan.routes:
            rows = [{"#": 0, "stop": f"Depot: {stops.name[0]}", "arrive": hhmm(r.start), "window": "", "units": "", "map": ""}]
            for seq, (n, t) in enumerate(zip(r.nodes, r.arrivals), start=1):
                rows.append({
                    "#": seq, "stop": stops.name[n], "arrive": hhmm(t),
                    "window": f"{hhmm(stops.tw_start[n])}–{hhmm(stops.tw_end[n])}", "units": int(stops.demand[n]),
                    "map": f"https://www.google.com/maps/search/?api=1&query={stops.lat[n]:.5f},{stops.lon[n]:.5f}",
                })
            rows.append({"#": "", "stop": "Back at depot", "arrive": hhmm(r.end), "window": "", "units": r.load, "map": ""})
            name = f"Van {r.vehicle}"
            pd.DataFrame(rows).to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            _style(ws, {1: 5, 2: 28, 3: 9, 4: 14, 5: 7, 6: 16})
            for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
                for cell in row:
                    if cell.value:
                        cell.hyperlink, cell.value, cell.font = cell.value, "Open map", Font(color="1F5FAD", underline="single")
            for i, link in enumerate(maps_legs(stops, r), start=1):
                cell = ws.cell(row=ws.max_row + 2, column=2, value=f"Navigation, part {i}")
                cell.hyperlink, cell.font = link, Font(color="1F5FAD", underline="single", bold=True)
    return buf.getvalue()
